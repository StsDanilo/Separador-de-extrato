from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = ROOT.parent
OUTPUT_DIR = ROOT / "validation_output"
CATEGORIES_PATH = ROOT / "config" / "categorias.json"

sys.path.insert(0, str(ROOT))

from separador_unificado.processors import banco_brasil, mercado_pago, stone  # noqa: E402
from separador_unificado.processors.detection import BANK_BB, BANK_MP, BANK_STONE, detect_bank  # noqa: E402


def is_generated_file(path: Path) -> bool:
    name = path.name.lower()
    return (
        "organizado" in name
        or "categorizado" in name
        or name.startswith("extrato_")
        or name.endswith("saida.xlsx")
    )


def source_files() -> list[Path]:
    roots = [
        WORKSPACE / "Separador extrato BB",
        WORKSPACE / "Separador extrato MP",
    ]
    files: list[Path] = []
    for root in roots:
        for path in root.rglob("*"):
            if path.suffix.lower() in {".xlsx", ".xlsm"} and not is_generated_file(path):
                files.append(path)
    stone_file = WORKSPACE / "extratoStoneJulho.xlsx"
    if stone_file.exists():
        files.append(stone_file)
    return sorted(files, key=lambda item: str(item).lower())


def verify_workbook(path: Path) -> tuple[str, int, int]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        return ws.title, ws.max_row, ws.max_column
    finally:
        wb.close()


def main() -> int:
    if OUTPUT_DIR.exists():
        shutil.rmtree(OUTPUT_DIR)
    OUTPUT_DIR.mkdir(parents=True)

    files = source_files()
    failures = []

    print(f"Arquivos candidatos: {len(files)}")
    for source in files:
        relative = source.relative_to(WORKSPACE)
        detected = detect_bank(source)
        print(f"\n- {relative}")
        print(f"  banco detectado: {detected or 'nao identificado'}")

        try:
            if detected == BANK_BB:
                result = banco_brasil.process_file(source, OUTPUT_DIR, CATEGORIES_PATH)
            elif detected == BANK_MP:
                result = mercado_pago.process_file(source, OUTPUT_DIR)
            elif detected == BANK_STONE:
                result = stone.process_file(source, OUTPUT_DIR)
            else:
                raise ValueError("Banco nao identificado")

            sheet, rows, cols = verify_workbook(result.output_path)
            print(f"  saida: {result.output_path.name}")
            print(f"  planilha: {sheet} | linhas: {rows} | colunas: {cols}")
            if result.banco == BANK_BB:
                print(f"  lancamentos: {result.total_rows} | entradas: {result.entradas} | saidas: {result.saidas}")
            elif result.banco == BANK_STONE:
                print(f"  lancamentos: {result.total_rows} | creditos: {result.entradas} | debitos: {result.saidas}")
            else:
                print(f"  lancamentos: {result.total_rows}")
        except Exception as exc:
            failures.append((relative, exc))
            print(f"  ERRO: {exc}")

    print("\nResumo")
    print(f"  processados: {len(files) - len(failures)}")
    print(f"  falhas: {len(failures)}")

    if failures:
        for relative, exc in failures:
            print(f"  - {relative}: {exc}")
        return 1

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
