from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


BANK_BB = "Banco do Brasil"
BANK_MP = "Mercado Pago"
BANK_STONE = "Stone"
SUPPORTED_BANKS = (BANK_BB, BANK_MP, BANK_STONE)


def _sem_acento(texto) -> str:
    return (
        unicodedata.normalize("NFKD", str(texto))
        .encode("ascii", "ignore")
        .decode("ascii")
        .lower()
    )


def _normalizar_rotulo(texto) -> str:
    texto = _sem_acento(texto)
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def detect_bank(path: str | Path) -> str | None:
    """Tenta identificar o banco pela estrutura do arquivo Excel."""
    workbook_path = Path(path)
    if not workbook_path.exists() or workbook_path.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
        return None

    if workbook_path.suffix.lower() == ".xls":
        return _detect_legacy_xls(workbook_path)

    try:
        wb = load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception:
        return None

    try:
        ws = wb.active
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True):
            rows.append([value for value in row if value not in (None, "")])

        return _detect_from_rows(rows)
    finally:
        wb.close()


def _detect_legacy_xls(path: Path) -> str | None:
    try:
        import xlrd
    except ModuleNotFoundError:
        return None

    try:
        wb = xlrd.open_workbook(str(path))
        sheet = wb.sheet_by_index(0)
        rows = []
        for row_idx in range(min(sheet.nrows, 12)):
            row = sheet.row_values(row_idx)
            rows.append([value for value in row if value not in (None, "")])
        return _detect_from_rows(rows)
    except Exception:
        return None


def _detect_from_rows(rows: list[list]) -> str | None:
    flat_raw = {str(value).strip() for row in rows for value in row}
    mp_headers = {"RELEASE_DATE", "TRANSACTION_TYPE", "TRANSACTION_NET_AMOUNT"}
    if len(mp_headers.intersection(flat_raw)) >= 2:
        return BANK_MP

    mp_summary_headers = {"INITIAL_BALANCE", "CREDITS", "DEBITS", "FINAL_BALANCE"}
    mp_translated_headers = {"Data", "Descricao", "Referencia", "Valor", "Saldo parcial"}
    if len(mp_summary_headers.intersection(flat_raw)) >= 2 and len(mp_translated_headers.intersection(flat_raw)) >= 3:
        return BANK_MP

    normalized = [_normalizar_rotulo(value) for row in rows for value in row]

    stone_headers = {
        "movimentaao",
        "movimentacao",
        "saldo antes",
        "saldo depois",
        "destino documento",
        "origem documento",
        "situacao",
    }
    if len(stone_headers.intersection(normalized)) >= 4:
        return BANK_STONE

    has_date = "data" in normalized or "dt" in normalized
    has_value = any(label == "valor" or label.startswith("valor") for label in normalized)
    has_history = any(label in {"historico", "lancamento", "movimento"} for label in normalized)
    has_type = any(label in {"inf", "info", "tipo", "tipo lancamento"} for label in normalized)
    if has_date and has_value and (has_history or has_type):
        return BANK_BB

    return None
