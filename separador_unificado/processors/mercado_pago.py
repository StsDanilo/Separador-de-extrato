from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .detection import BANK_MP
from .models import ProcessResult


HEADER_ROWS = 4
DATE_COLUMN = 1
AMOUNT_COLUMN = 4
DESCRIPTION_COLUMN = 2

TRANSLATED_HEADERS = {
    "RELEASE_DATE": "Data",
    "TRANSACTION_TYPE": "Descricao",
    "REFERENCE_ID": "Referencia",
    "TRANSACTION_NET_AMOUNT": "Valor",
    "PARTIAL_BALANCE": "Saldo parcial",
}

DAY_FILL_LIGHT = "F2F4F7"
DAY_FILL_WHITE = "FFFFFF"
TITLE_FILL = "1F4E78"
TABLE_HEADER_FILL = "305496"
BORDER_COLOR = "D9E2F3"


@dataclass(frozen=True)
class TransactionRow:
    values: list
    source_cells: list[Cell]
    original_index: int
    date: datetime
    amount: float


def parse_brazilian_number(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    if not text:
        return 0.0

    text = text.replace(".", "").replace(",", ".")
    return float(text)


def parse_date(value) -> datetime:
    if isinstance(value, datetime):
        return value

    text = str(value).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue

    raise ValueError(f"Data invalida no extrato: {value!r}")


def copy_cell_style(source: Cell, target: Cell) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def read_transactions(ws) -> list[TransactionRow]:
    transactions: list[TransactionRow] = []

    for row_number in range(HEADER_ROWS + 1, ws.max_row + 1):
        source_cells = list(ws[row_number])
        values = [cell.value for cell in source_cells]

        if all(value in (None, "") for value in values):
            continue

        date = parse_date(values[DATE_COLUMN - 1])
        amount = parse_brazilian_number(values[AMOUNT_COLUMN - 1])

        transactions.append(
            TransactionRow(
                values=values,
                source_cells=source_cells,
                original_index=row_number,
                date=date,
                amount=amount,
            )
        )

    return transactions


def organize_transactions(transactions: Iterable[TransactionRow]) -> list[TransactionRow]:
    by_day: dict[datetime, list[TransactionRow]] = {}
    for transaction in transactions:
        day = transaction.date.replace(hour=0, minute=0, second=0, microsecond=0)
        by_day.setdefault(day, []).append(transaction)

    organized: list[TransactionRow] = []
    for day_index, day in enumerate(sorted(by_day)):
        rows = by_day[day]
        positives = [row for row in rows if row.amount >= 0]
        negatives = [row for row in rows if row.amount < 0]

        if day_index % 2 == 0:
            organized.extend(positives + negatives)
        else:
            organized.extend(negatives + positives)

    return organized


def build_output(input_path: Path, output_path: Path) -> int:
    source_wb = load_workbook(input_path)
    source_ws = source_wb.active

    output_wb = Workbook()
    output_ws = output_wb.active
    output_ws.title = "Extrato organizado"

    max_col = source_ws.max_column

    for row_number in range(1, HEADER_ROWS + 1):
        for col_number in range(1, max_col + 1):
            source_cell = source_ws.cell(row=row_number, column=col_number)
            value = TRANSLATED_HEADERS.get(source_cell.value, source_cell.value)
            target_cell = output_ws.cell(row=row_number, column=col_number, value=value)
            copy_cell_style(source_cell, target_cell)

    transactions = organize_transactions(read_transactions(source_ws))

    thin_border = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="thin", color=BORDER_COLOR),
        bottom=Side(style="thin", color=BORDER_COLOR),
    )

    current_day = None
    day_index = -1

    for output_row, transaction in enumerate(transactions, start=HEADER_ROWS + 1):
        transaction_day = transaction.date.date()
        if transaction_day != current_day:
            current_day = transaction_day
            day_index += 1

        fill_color = DAY_FILL_LIGHT if day_index % 2 == 0 else DAY_FILL_WHITE

        for col_number, value in enumerate(transaction.values[:max_col], start=1):
            source_cell = transaction.source_cells[col_number - 1]
            target_cell = output_ws.cell(row=output_row, column=col_number, value=value)
            copy_cell_style(source_cell, target_cell)
            target_cell.fill = PatternFill("solid", fgColor=fill_color)
            target_cell.border = thin_border
            target_cell.alignment = Alignment(vertical="center")
            if col_number == DESCRIPTION_COLUMN:
                target_cell.alignment = Alignment(vertical="center", wrap_text=True)

    style_headers(output_ws, max_col, len(transactions))
    resize_columns(output_ws, max_col)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(output_path)
    source_wb.close()
    output_wb.close()
    return len(transactions)


def style_headers(ws, max_col: int, transaction_count: int) -> None:
    last_row = HEADER_ROWS + transaction_count

    summary_header = ws[1]
    summary_values = ws[2]
    table_header = ws[HEADER_ROWS]

    for cell in summary_header[:max_col]:
        cell.fill = PatternFill("solid", fgColor=TITLE_FILL)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for cell in summary_values[:max_col]:
        cell.font = Font(bold=True, color="1F2937")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for cell in table_header[:max_col]:
        cell.fill = PatternFill("solid", fgColor=TABLE_HEADER_FILL)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[4].height = 24
    ws.freeze_panes = "A5"

    if last_row >= HEADER_ROWS:
        ws.auto_filter.ref = f"A{HEADER_ROWS}:{get_column_letter(max_col)}{last_row}"


def resize_columns(ws, max_col: int) -> None:
    min_widths = {
        1: 10,
        2: 48,
        3: 18,
        4: 22,
        5: 18,
    }

    for col_number in range(1, max_col + 1):
        letter = get_column_letter(col_number)
        max_length = 0
        for cell in ws[letter]:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        if col_number == DATE_COLUMN:
            width = 10
        else:
            width = min(max(max_length + 2, min_widths.get(col_number, 12)), 60)
        ws.column_dimensions[letter].width = width


def default_output_path(input_path: Path, output_dir: Path | None = None) -> Path:
    generated_at = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{input_path.stem}_organizado_{generated_at}{input_path.suffix}"
    return (output_dir or input_path.parent) / filename


def process_file(input_path: str | Path, output_dir: str | Path | None = None) -> ProcessResult:
    input_path = Path(input_path)
    output_root = Path(output_dir) if output_dir else None

    if not input_path.exists():
        raise FileNotFoundError(f"Arquivo de entrada nao encontrado: {input_path}")

    output_path = default_output_path(input_path, output_root)
    total = build_output(input_path, output_path)
    return ProcessResult(output_path=output_path, total_rows=total, banco=BANK_MP)

