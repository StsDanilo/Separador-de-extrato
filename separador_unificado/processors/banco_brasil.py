from __future__ import annotations

import json
import re
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .detection import BANK_BB
from .models import ProcessResult


CAMPO_SAIDA = [
    ("data", "Data"),
    ("historico", "Historico"),
    ("detalhes", "Detalhes"),
    ("documento", "No documento"),
    ("valor", "Valor"),
    ("tipo", "Tipo"),
]


def carregar_categorias(path: str | Path | None) -> dict:
    if not path:
        return {}

    categorias_path = Path(path)
    if not categorias_path.exists():
        return {}

    with categorias_path.open(encoding="utf-8") as f:
        data = json.load(f)

    if not isinstance(data, dict):
        raise ValueError("O arquivo de categorias precisa conter um objeto JSON.")
    return data


def sem_acento(texto) -> str:
    return (
        unicodedata.normalize("NFKD", str(texto))
        .encode("ascii", "ignore")
        .decode("ascii")
        .lower()
    )


def normalizar_rotulo(texto) -> str:
    texto = sem_acento(texto)
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def normalizar_valor(texto) -> str:
    return str(texto).strip() if texto is not None else ""


def normalizar_valor_monetario(texto) -> str:
    valor = normalizar_valor(texto)
    if not valor:
        return ""

    valor = re.sub(r"\s+[cd]$", "", valor, flags=re.IGNORECASE).strip()
    valor = re.sub(r"^-+\s*", "", valor).strip()
    return valor


def normalizar_tipo(tipo) -> str:
    t = sem_acento(tipo).replace(" ", "").strip()
    if not t:
        return ""
    if t.startswith(("c", "entrada")):
        return "c"
    if t.startswith(("d", "saida")):
        return "d"
    if t.startswith("sa") and t.endswith("da"):
        return "d"
    return t[:1]


def eh_entrada(tipo) -> bool:
    return normalizar_tipo(tipo) == "c"


def eh_saida(tipo) -> bool:
    return normalizar_tipo(tipo) == "d"


def _match_header(rotulo, opcoes_exatas=None, opcoes_inicio=None, opcoes_contem=None) -> bool:
    opcoes_exatas = opcoes_exatas or []
    opcoes_inicio = opcoes_inicio or []
    opcoes_contem = opcoes_contem or []

    if rotulo in opcoes_exatas:
        return True
    if any(rotulo.startswith(op) for op in opcoes_inicio):
        return True
    if any(op in rotulo for op in opcoes_contem):
        return True
    return False


def identificar_colunas(ws, max_linhas_busca: int = 10):
    melhor = ({}, -1, None)

    limite = min(ws.max_row, max_linhas_busca)
    for r in range(1, limite + 1):
        row = [cell.value for cell in ws[r]]
        col_map = {}
        for i, valor in enumerate(row):
            if valor is None:
                continue
            rotulo = normalizar_rotulo(valor)
            if not rotulo:
                continue

            if "data" not in col_map and _match_header(rotulo, opcoes_exatas=["data", "dt"]):
                col_map["data"] = i
            elif "historico" not in col_map and _match_header(
                rotulo, opcoes_exatas=["historico", "lancamento", "movimento"]
            ):
                col_map["historico"] = i
            elif "detalhes" not in col_map and _match_header(
                rotulo,
                opcoes_exatas=["detalhes", "detalhamento", "complemento", "info complementar"],
                opcoes_inicio=["detalhamento hist"],
                opcoes_contem=["detalhamento hist"],
            ):
                col_map["detalhes"] = i
            elif "documento" not in col_map and _match_header(
                rotulo,
                opcoes_exatas=[
                    "numero documento",
                    "num documento",
                    "n documento",
                    "documento",
                    "doc",
                    "n doc",
                    "no documento",
                ],
                opcoes_inicio=["numero documento", "num documento", "n documento", "no documento"],
                opcoes_contem=["numero documento"],
            ):
                col_map["documento"] = i
            elif "valor" not in col_map and _match_header(
                rotulo, opcoes_exatas=["valor r", "valor", "vlr"], opcoes_inicio=["valor"]
            ):
                col_map["valor"] = i
            elif "tipo" not in col_map and _match_header(
                rotulo,
                opcoes_exatas=["inf", "info", "tipo", "tipo lancamento"],
                opcoes_inicio=["tipo lancamento"],
            ):
                col_map["tipo"] = i

        score = len(col_map)
        if score > melhor[1]:
            melhor = (col_map, score, r)

        if score >= 5 and "data" in col_map and "historico" in col_map and "valor" in col_map:
            return col_map, r

    if melhor[0]:
        return melhor[0], melhor[2]

    raise ValueError("Nao foi possivel identificar o cabecalho do extrato.")


def montar_linha_saida(row, col_map):
    linha = []
    for campo, _titulo in CAMPO_SAIDA:
        idx = col_map.get(campo)
        valor = row[idx] if idx is not None and idx < len(row) else ""
        if campo == "tipo":
            tipo = normalizar_tipo(valor)
            if tipo == "c":
                valor = "(+)"
            elif tipo == "d":
                valor = "(-)"
            else:
                valor = normalizar_valor(valor)
        elif campo == "valor":
            valor = normalizar_valor_monetario(valor)
        else:
            valor = normalizar_valor(valor)
        linha.append(valor)
    return linha


def linha_corresponde(row_data, condicoes) -> bool:
    lancamento = row_data.get("historico", "") or row_data.get("lancamento", "")
    detalhes = row_data.get("detalhes", "")

    for cond in condicoes:
        campo = cond.get("campo", "")
        termos = cond.get("termos", [])
        modo = cond.get("modo", "contem")

        if campo in ("lancamento", "historico"):
            texto = lancamento
        elif campo in ("detalhes", "detalhamento"):
            texto = detalhes
        else:
            texto = f"{lancamento} {detalhes}"

        texto_norm = sem_acento(texto)
        if modo == "contem":
            if not any(sem_acento(t) in texto_norm for t in termos):
                return False
        elif modo == "exato":
            if not any(sem_acento(t) == texto_norm for t in termos):
                return False
        else:
            raise ValueError(f"Modo de categoria invalido: {modo}")

    return True


def categorizar_linha(row_data, categorias) -> str:
    for nome, cat in categorias.items():
        condicoes = cat.get("condicoes", [])
        if linha_corresponde(row_data, condicoes):
            return nome
    return "Outros"


def eh_linha_vazia(row) -> bool:
    return all(v is None or str(v).strip() == "" for v in row)


def ler_extrato(path: str | Path):
    wb = load_workbook(path)
    ws = wb.active

    col_map, linha_cabecalho = identificar_colunas(ws)
    headers_saida = [titulo for _campo, titulo in CAMPO_SAIDA]

    entradas = []
    saidas = []
    saldo_anterior = None

    for row in ws.iter_rows(min_row=linha_cabecalho + 1, values_only=True):
        row = list(row)
        if eh_linha_vazia(row):
            continue

        row_saida = montar_linha_saida(row, col_map)
        historico = row_saida[1]
        historico_norm = re.sub(r"\s+", " ", sem_acento(historico)).strip()

        if "saldo anterior" in historico_norm:
            saldo_anterior = row_saida
            continue

        if historico_norm == "s a l d o" or "saldo" in historico_norm:
            continue

        idx_tipo = col_map.get("tipo")
        tipo_bruto = row[idx_tipo] if idx_tipo is not None and idx_tipo < len(row) else ""
        if eh_entrada(tipo_bruto):
            entradas.append(row_saida)
        elif eh_saida(tipo_bruto):
            saidas.append(row_saida)

    wb.close()
    return headers_saida, col_map, saldo_anterior, entradas, saidas


def agrupar_por_categoria(linhas, categorias):
    grupos = {nome: [] for nome in categorias}
    grupos["Outros"] = []

    for row in linhas:
        row_data = {
            "historico": row[1] if len(row) > 1 else "",
            "lancamento": row[1] if len(row) > 1 else "",
            "detalhes": row[2] if len(row) > 2 else "",
        }
        cat = categorizar_linha(row_data, categorias)
        grupos[cat].append(row)

    return grupos


FILL_DIA_A = PatternFill("solid", fgColor="FFFFFF")
FILL_DIA_B = PatternFill("solid", fgColor="EFEFEF")
FILL_BRANCO = PatternFill("solid", fgColor="FFFFFF")


def borda_fina():
    s = Side(style="thin", color="000000")
    return Border(left=s, right=s, top=s, bottom=s)


def _aplicar_merged_row(ws, row_idx, num_cols, value, font, alignment, fill):
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=num_cols)
    borda = borda_fina()
    borda_interna = Border(top=borda.top, bottom=borda.bottom)
    for c in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = fill
        cell.border = borda_interna

    ws.cell(row=row_idx, column=1).border = Border(
        left=borda.left, right=Side(style=None), top=borda.top, bottom=borda.bottom
    )
    ws.cell(row=row_idx, column=num_cols).border = Border(
        left=Side(style=None), right=borda.right, top=borda.top, bottom=borda.bottom
    )
    master = ws.cell(row=row_idx, column=1)
    master.value = value
    master.font = font
    master.alignment = alignment


def escrever_header_bloco(ws, row_idx, titulo, num_cols):
    _aplicar_merged_row(
        ws,
        row_idx,
        num_cols,
        titulo,
        Font(bold=True, size=13),
        Alignment(horizontal="center", vertical="center"),
        FILL_BRANCO,
    )
    ws.row_dimensions[row_idx].height = 22


def escrever_header_categoria(ws, row_idx, nome, num_cols):
    _aplicar_merged_row(
        ws,
        row_idx,
        num_cols,
        f"  {nome}",
        Font(bold=True, size=10),
        Alignment(horizontal="left", vertical="center"),
        FILL_BRANCO,
    )
    ws.row_dimensions[row_idx].height = 18


def escrever_header_colunas(ws, row_idx, headers):
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=c_idx, value=h)
        cell.font = Font(bold=True, size=9)
        cell.alignment = Alignment(horizontal="center")
        cell.border = borda_fina()
        cell.fill = FILL_BRANCO


def escrever_linha_dados(ws, row_idx, row, fill=None):
    borda = borda_fina()
    for c_idx, val in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=c_idx, value=val)
        cell.font = Font(size=9)
        cell.border = borda
        if fill:
            cell.fill = fill

        if c_idx == 1:
            cell.alignment = Alignment(horizontal="center")
        elif c_idx == 5:
            cell.alignment = Alignment(horizontal="right")
        elif c_idx in (2, 3, 4):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        elif c_idx == 6:
            cell.alignment = Alignment(horizontal="center")


def linha_em_branco(ws, row_idx, num_cols):
    for c in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.value = None
        cell.fill = PatternFill(fill_type=None)
        cell.border = Border()


def escrever_bloco(ws, current_row, titulo, grupos, headers, num_cols, sufixo_categoria=""):
    escrever_header_bloco(ws, current_row, titulo, num_cols)
    current_row += 1

    idx_data = 0
    primeiro = True
    for nome_cat, linhas in grupos.items():
        if not linhas:
            continue

        if not primeiro:
            linha_em_branco(ws, current_row, num_cols)
            current_row += 1

        nome_exibido = f"{nome_cat}{sufixo_categoria}" if sufixo_categoria else nome_cat
        escrever_header_categoria(ws, current_row, nome_exibido, num_cols)
        current_row += 1

        escrever_header_colunas(ws, current_row, headers)
        current_row += 1

        data_atual = None
        fill_par = False
        for row in linhas:
            data_linha = row[idx_data] if idx_data < len(row) else ""
            if data_linha != data_atual:
                data_atual = data_linha
                fill_par = not fill_par
            fill = FILL_DIA_A if fill_par else FILL_DIA_B
            escrever_linha_dados(ws, current_row, row, fill=fill)
            current_row += 1

        primeiro = False

    return current_row


def default_output_path(input_path: Path, output_dir: Path | None = None) -> Path:
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    filename = f"{input_path.stem}_categorizado_{timestamp}.xlsx"
    return (output_dir or input_path.parent) / filename


def gerar_extrato(
    input_path: str | Path,
    output_path: str | Path,
    categorias_path: str | Path | None = None,
) -> ProcessResult:
    categorias = carregar_categorias(categorias_path)
    headers, _col_map, saldo_anterior, entradas, saidas = ler_extrato(input_path)

    num_cols = len(headers)
    wb_out = Workbook()
    ws = wb_out.active
    ws.title = "Extrato Categorizado"

    current_row = 1

    if saldo_anterior:
        escrever_linha_dados(ws, current_row, saldo_anterior)
        ws.cell(row=current_row, column=1).font = Font(bold=True, size=9)
        current_row += 1

    linha_em_branco(ws, current_row, num_cols)
    current_row += 1

    grupos_entrada = agrupar_por_categoria(entradas, categorias)
    current_row = escrever_bloco(
        ws, current_row, "ENTRADAS", grupos_entrada, headers, num_cols, sufixo_categoria="-E"
    )

    linha_em_branco(ws, current_row, num_cols)
    current_row += 1
    linha_em_branco(ws, current_row, num_cols)
    current_row += 1

    grupos_saida = agrupar_por_categoria(saidas, categorias)
    escrever_bloco(ws, current_row, "SAÍDAS", grupos_saida, headers, num_cols, sufixo_categoria="-D")

    larguras_fixas = {
        1: 11,
        2: 22,
        3: 28,
        4: 16,
        5: 12,
        6: 10,
    }
    for c_idx in range(1, num_cols + 1):
        col_letter = get_column_letter(c_idx)
        if c_idx in larguras_fixas:
            ws.column_dimensions[col_letter].width = larguras_fixas[c_idx]
        else:
            max_len = 0
            for cell in ws[col_letter]:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(output_path)
    wb_out.close()

    return ProcessResult(
        output_path=output_path,
        total_rows=len(entradas) + len(saidas),
        entradas=len(entradas),
        saidas=len(saidas),
        banco=BANK_BB,
    )


def process_file(
    input_path: str | Path,
    output_dir: str | Path | None = None,
    categorias_path: str | Path | None = None,
) -> ProcessResult:
    input_path = Path(input_path)
    if not input_path.exists():
        raise FileNotFoundError(f"Arquivo de entrada nao encontrado: {input_path}")

    output_root = Path(output_dir) if output_dir else None
    output_path = default_output_path(input_path, output_root)
    return gerar_extrato(input_path, output_path, categorias_path)
