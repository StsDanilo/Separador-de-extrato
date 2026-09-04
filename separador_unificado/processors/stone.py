from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field as dataclass_field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .detection import BANK_STONE
from .models import ProcessResult


HEADER_ALIASES = {
    "movimentacao": {"movimentacao", "movimentaao"},
    "tipo": {"tipo"},
    "valor": {"valor"},
    "saldo_antes": {"saldo antes"},
    "saldo_depois": {"saldo depois"},
    "tarifa": {"tarifa"},
    "data": {"data"},
    "nosso_numero": {"nosso numero", "nosso nmero", "nosso numero"},
    "situacao": {"situacao", "situaao"},
    "destino": {"destino"},
    "destino_documento": {"destino documento"},
    "destino_instituicao": {"destino instituicao", "destino instituiao"},
    "destino_agencia": {"destino agencia", "destino agncia"},
    "destino_conta": {"destino conta"},
    "origem": {"origem"},
    "origem_documento": {"origem documento"},
    "origem_instituicao": {"origem instituicao", "origem instituiao"},
    "origem_agencia": {"origem agencia", "origem agncia"},
    "origem_conta": {"origem conta"},
}

DISPLAY_HEADERS = {
    "movimentacao": "Movimentação",
    "tipo": "Tipo",
    "valor": "Valor",
    "saldo_antes": "Saldo antes",
    "saldo_depois": "Saldo depois",
    "tarifa": "Tarifa",
    "data": "Data",
    "nosso_numero": "Nosso Número",
    "situacao": "Situação",
    "destino": "Destino",
    "destino_documento": "Destino Documento",
    "destino_instituicao": "Destino Instituição",
    "destino_agencia": "Destino Agência",
    "destino_conta": "Destino Conta",
    "origem": "Origem",
    "origem_documento": "Origem Documento",
    "origem_instituicao": "Origem Instituição",
    "origem_agencia": "Origem Agência",
    "origem_conta": "Origem Conta",
}

OUTPUT_ORDER = [
    "destino",
    "data",
    "valor",
    "saldo_depois",
]

ALIGN_CENTER_FIELDS = {"data"}
ALIGN_RIGHT_FIELDS = {"valor", "saldo_antes", "saldo_depois"}

FIELD_WIDTHS = {
    "data": 17,
    "movimentacao": 15,
    "valor": 14,
    "destino": 32,
    "origem": 32,
    "saldo_antes": 16,
    "saldo_depois": 16,
}

FILL_DIA_A = PatternFill("solid", fgColor="FFFFFF")
FILL_DIA_B = PatternFill("solid", fgColor="EFEFEF")
FILL_HEADER = PatternFill("solid", fgColor="FFFFFF")
FILL_GROUP = PatternFill("solid", fgColor="F2F4F7")
FILL_SECTION = PatternFill("solid", fgColor="E6EAEE")
FILL_INFO = PatternFill("solid", fgColor="F7F9FB")


def sem_acento(texto) -> str:
    return (
        unicodedata.normalize("NFKD", str(texto))
        .encode("ascii", "ignore")
        .decode("ascii")
        .lower()
    )


def normalizar_rotulo(texto) -> str:
    texto = sem_acento(texto)
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def normalizar_valor(texto) -> str:
    return str(texto).strip() if texto is not None else ""


def normalizar_movimentacao(texto) -> str:
    valor = sem_acento(texto).replace(" ", "").strip()
    if valor.startswith("cr"):
        return "credito"
    if valor.startswith("d"):
        return "debito"
    return valor


def eh_desconhecido(texto) -> bool:
    return normalizar_rotulo(texto) == "desconhecido"


def _load_xlsx_rows(path: Path) -> list[list]:
    # Não usar read_only=True: alguns extratos da Stone trazem a tag <dimension>
    # do XML incorreta (ex.: "A1" em vez do intervalo real), e nesse modo o
    # openpyxl confia nela para decidir até onde ler, cortando quase todo o
    # arquivo. Carregar normalmente força a varredura das células reais.
    wb = load_workbook(path, data_only=True)
    try:
        ws = wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _load_xls_rows(path: Path) -> list[list]:
    try:
        import xlrd
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Para ler arquivos .xls da Stone, instale as dependencias com: "
            "python -m pip install -r requirements.txt"
        ) from exc

    try:
        workbook = xlrd.open_workbook(str(path))
    except Exception as exc:
        raise RuntimeError(
            "Nao consegui abrir este .xls como planilha Excel antiga. "
            "Se o arquivo foi exportado pelo banco nesse formato, abra no Excel/LibreOffice "
            "e salve como .xlsx antes de processar."
        ) from exc
    sheet = workbook.sheet_by_index(0)
    rows: list[list] = []
    for row_idx in range(sheet.nrows):
        values = []
        for col_idx in range(sheet.ncols):
            cell = sheet.cell(row_idx, col_idx)
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                value = xlrd.xldate_as_datetime(value, workbook.datemode)
            values.append(value)
        rows.append(values)
    return rows


def load_rows(path: str | Path) -> list[list]:
    input_path = Path(path)
    suffix = input_path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _load_xlsx_rows(input_path)
    if suffix == ".xls":
        return _load_xls_rows(input_path)
    raise ValueError("A Stone aceita arquivos .xlsx, .xlsm ou .xls.")


# Colunas indispensaveis para aplicar as regras de separacao da Stone
# (resgate/aplicacao/comuns) e para o extrato fazer sentido. Sem qualquer uma
# delas nao da para dividir os lancamentos com seguranca.
CORE_HEADERS = ("movimentacao", "valor", "data", "destino", "origem")

_ALL_KNOWN_ALIASES = {alias for aliases in HEADER_ALIASES.values() for alias in aliases}


@dataclass
class HeaderMatch:
    header_idx: int
    col_map: dict[str, int]
    is_exact: bool
    # campo -> texto da coluna usada quando o nome nao bateu com nenhum
    # apelido conhecido, so uma aproximacao (ex.: "Conta Destino" para "destino")
    fuzzy_matches: dict[str, str] = dataclass_field(default_factory=dict)
    # nomes de exibicao das colunas essenciais que nao foram encontradas nem por aproximacao
    missing: list[str] = dataclass_field(default_factory=list)


def _fuzzy_match_column(field: str, row: list, used_cols: set[int]) -> tuple[int, str] | None:
    """Procura, entre as colunas ainda nao usadas, alguma cujo rotulo contenha
    a palavra do campo procurado (ex.: "Nome Destino" contem "destino")."""
    keywords = HEADER_ALIASES.get(field, {field})
    for col_idx, value in enumerate(row):
        if col_idx in used_cols:
            continue
        label = normalizar_rotulo(value)
        if not label:
            continue
        if label in _ALL_KNOWN_ALIASES:
            # Rotulo bate exatamente com outro campo conhecido; nao e uma
            # aproximacao do campo atual, e arriscar essa troca pode misturar
            # colunas com significados diferentes (ex.: "Origem Documento").
            continue
        if any(word in keywords for word in label.split()):
            return col_idx, str(value).strip()
    return None


def identify_header(rows: list[list]) -> HeaderMatch:
    best_row_idx = None
    best_map: dict[str, int] = {}
    best_core_matches = -1

    for row_idx, row in enumerate(rows[:15]):
        col_map: dict[str, int] = {}
        for col_idx, value in enumerate(row):
            label = normalizar_rotulo(value)
            if not label:
                continue
            for field, aliases in HEADER_ALIASES.items():
                if field not in col_map and label in aliases:
                    col_map[field] = col_idx
                    break

        core_matches = sum(1 for field in CORE_HEADERS if field in col_map)
        if core_matches == len(CORE_HEADERS):
            return HeaderMatch(row_idx, col_map, is_exact=True)

        is_better = core_matches > best_core_matches or (
            core_matches == best_core_matches and len(col_map) > len(best_map)
        )
        if is_better:
            best_row_idx = row_idx
            best_map = col_map
            best_core_matches = core_matches

    if best_row_idx is None:
        missing = [DISPLAY_HEADERS[field] for field in CORE_HEADERS]
        return HeaderMatch(-1, {}, is_exact=False, missing=missing)

    row = rows[best_row_idx]
    used_cols = set(best_map.values())
    fuzzy_matches: dict[str, str] = {}
    for field in CORE_HEADERS:
        if field in best_map:
            continue
        found = _fuzzy_match_column(field, row, used_cols)
        if found is None:
            continue
        col_idx, label_text = found
        best_map[field] = col_idx
        used_cols.add(col_idx)
        fuzzy_matches[field] = label_text

    missing = [DISPLAY_HEADERS[field] for field in CORE_HEADERS if field not in best_map]
    return HeaderMatch(best_row_idx, best_map, is_exact=False, fuzzy_matches=fuzzy_matches, missing=missing)


def format_missing_columns_message(missing: list[str]) -> str:
    if len(missing) == 1:
        return f'Não foi possível separar os lançamentos da Stone: não encontrei a coluna "{missing[0]}" no extrato.'
    colunas = ", ".join(f'"{nome}"' for nome in missing)
    return f"Não foi possível separar os lançamentos da Stone: não encontrei as colunas {colunas} no extrato."


def row_is_empty(row: list) -> bool:
    return all(value is None or str(value).strip() == "" for value in row)


def output_headers(col_map: dict[str, int]) -> list[str]:
    headers = []
    for field in OUTPUT_ORDER:
        if field in col_map:
            headers.append(DISPLAY_HEADERS[field])
    return headers


def output_fields(col_map: dict[str, int]) -> list[str]:
    return [field for field in OUTPUT_ORDER if field in col_map]


def output_row(row: list, col_map: dict[str, int]) -> list:
    values = []
    for field in OUTPUT_ORDER:
        if field not in col_map:
            continue
        idx = col_map[field]
        value = row[idx] if idx < len(row) else ""
        if isinstance(value, datetime):
            value = value.strftime("%d/%m/%Y %H:%M")
        values.append(normalizar_valor(value))
    return values


def classify_row(row: list, col_map: dict[str, int]) -> str:
    mov_idx = col_map.get("movimentacao")
    if mov_idx is None:
        # Sem a coluna Movimentacao nao ha como saber se e credito ou debito;
        # o lancamento ainda entra na planilha, só que no bloco "Outros".
        return "Outros"

    movimento = row[mov_idx] if mov_idx < len(row) else ""
    movimento_norm = normalizar_movimentacao(movimento)

    dest_idx = col_map.get("destino")
    orig_idx = col_map.get("origem")
    destino = row[dest_idx] if dest_idx is not None and dest_idx < len(row) else ""
    origem = row[orig_idx] if orig_idx is not None and orig_idx < len(row) else ""

    if movimento_norm == "credito" and orig_idx is not None and eh_desconhecido(origem):
        return "Resgate da Reserva Stone"
    if movimento_norm == "credito":
        return "Créditos comuns"
    if movimento_norm == "debito" and dest_idx is not None and eh_desconhecido(destino):
        return "Aplicação na Reserva Stone"
    if movimento_norm == "debito":
        return "Débitos comuns"
    return "Outros"


def parse_data_stone(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    text = str(value).strip() if value is not None else ""
    if not text:
        return None
    for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def ordenar_cronologicamente(rows: list[list], header_idx: int, col_map: dict[str, int]) -> list[list]:
    data_rows = [row for row in rows[header_idx + 1 :] if not row_is_empty(row)]
    if not data_rows:
        return data_rows

    data_idx = col_map.get("data")
    if data_idx is None:
        return data_rows

    primeira_data = parse_data_stone(data_rows[0][data_idx] if data_idx < len(data_rows[0]) else None)
    ultima_data = parse_data_stone(data_rows[-1][data_idx] if data_idx < len(data_rows[-1]) else None)
    if primeira_data is not None and ultima_data is not None and primeira_data > ultima_data:
        return list(reversed(data_rows))

    return data_rows


def compute_balance_summary(
    data_rows: list[list], col_map: dict[str, int]
) -> tuple[str | None, str | None]:
    saldo_antes_idx = col_map.get("saldo_antes")
    saldo_depois_idx = col_map.get("saldo_depois")
    if saldo_antes_idx is None or saldo_depois_idx is None or not data_rows:
        return None, None

    primeiro_lancamento = data_rows[0]
    ultimo_lancamento = data_rows[-1]

    saldo_inicial = None
    if saldo_antes_idx < len(primeiro_lancamento):
        saldo_inicial = normalizar_valor(primeiro_lancamento[saldo_antes_idx]) or None

    saldo_final = None
    if saldo_depois_idx < len(ultimo_lancamento):
        saldo_final = normalizar_valor(ultimo_lancamento[saldo_depois_idx]) or None

    return saldo_inicial, saldo_final


def split_rows(data_rows: list[list], col_map: dict[str, int]):
    groups = {
        "Resgate da Reserva Stone": [],
        "Créditos comuns": [],
        "Aplicação na Reserva Stone": [],
        "Débitos comuns": [],
        "Outros": [],
    }

    for row in data_rows:
        group = classify_row(row, col_map)
        groups[group].append(output_row(row, col_map))

    return groups


def borda_fina():
    side = Side(style="thin", color="000000")
    return Border(left=side, right=side, top=side, bottom=side)


def aplicar_linha_mesclada(ws, row_idx: int, num_cols: int, text: str, font: Font, fill: PatternFill):
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=num_cols)
    border = borda_fina()
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.fill = fill
        cell.border = Border(top=border.top, bottom=border.bottom)
    ws.cell(row=row_idx, column=1).border = Border(
        left=border.left, top=border.top, bottom=border.bottom
    )
    ws.cell(row=row_idx, column=num_cols).border = Border(
        right=border.right, top=border.top, bottom=border.bottom
    )
    master = ws.cell(row=row_idx, column=1)
    master.value = text
    master.font = font
    master.alignment = Alignment(horizontal="center", vertical="center")


def escrever_header_secao(ws, row_idx: int, titulo: str, num_cols: int):
    aplicar_linha_mesclada(ws, row_idx, num_cols, titulo, Font(bold=True, size=13), FILL_SECTION)
    ws.row_dimensions[row_idx].height = 23


def escrever_header_grupo(ws, row_idx: int, titulo: str, num_cols: int):
    aplicar_linha_mesclada(ws, row_idx, num_cols, f"  {titulo}", Font(bold=True, size=10), FILL_GROUP)
    ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row_idx].height = 19


def escrever_header_colunas(ws, row_idx: int, headers: list[str]):
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=header)
        cell.fill = FILL_HEADER
        cell.border = borda_fina()
        cell.font = Font(bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def escrever_linha(ws, row_idx: int, values: list, fill: PatternFill, fields: list[str]):
    border = borda_fina()
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.fill = fill
        cell.border = border
        cell.font = Font(size=9)
        field = fields[col_idx - 1] if col_idx - 1 < len(fields) else None
        if field in ALIGN_RIGHT_FIELDS:
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif field in ALIGN_CENTER_FIELDS:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def linha_em_branco(ws, row_idx: int, num_cols: int):
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.value = None
        cell.border = Border()
        cell.fill = PatternFill(fill_type=None)


def escrever_grupo(
    ws, current_row: int, group_name: str, rows: list[list], headers: list[str], fields: list[str], num_cols: int
):
    if not rows:
        return current_row

    escrever_header_grupo(ws, current_row, group_name, num_cols)
    current_row += 1
    escrever_header_colunas(ws, current_row, headers)
    current_row += 1

    data_col = fields.index("data") if "data" in fields else None

    current_day = None
    fill_toggle = False
    for values in rows:
        if data_col is not None and data_col < len(values):
            day = str(values[data_col]).split(" ")[0]
        else:
            day = ""
        if day != current_day:
            current_day = day
            fill_toggle = not fill_toggle
        escrever_linha(ws, current_row, values, FILL_DIA_A if fill_toggle else FILL_DIA_B, fields)
        current_row += 1

    return current_row


def escrever_secao(
    ws,
    current_row: int,
    title: str,
    group_names: list[str],
    groups: dict[str, list[list]],
    headers: list[str],
    fields: list[str],
    num_cols: int,
):
    escrever_header_secao(ws, current_row, title, num_cols)
    current_row += 1

    first = True
    for group_name in group_names:
        rows = groups.get(group_name, [])
        if not rows:
            continue
        if not first:
            linha_em_branco(ws, current_row, num_cols)
            current_row += 1
        current_row = escrever_grupo(ws, current_row, group_name, rows, headers, fields, num_cols)
        first = False

    return current_row


def escrever_linha_info(ws, row_idx: int, texto: str, num_cols: int):
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=num_cols)
    border = borda_fina()
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.fill = FILL_INFO
        cell.border = Border(top=border.top, bottom=border.bottom)
    ws.cell(row=row_idx, column=1).border = Border(left=border.left, top=border.top, bottom=border.bottom)
    ws.cell(row=row_idx, column=num_cols).border = Border(right=border.right, top=border.top, bottom=border.bottom)
    master = ws.cell(row=row_idx, column=1)
    master.value = texto
    master.font = Font(bold=True, size=11)
    master.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row_idx].height = 20


def ajustar_colunas(ws, fields: list[str]):
    for col_idx, field in enumerate(fields, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = FIELD_WIDTHS.get(field, 16)


def default_output_path(input_path: Path, output_dir: Path | None = None) -> Path:
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    filename = f"{input_path.stem}_stone_organizado_{timestamp}.xlsx"
    return (output_dir or input_path.parent) / filename


def analyze_columns(input_path: str | Path) -> HeaderMatch:
    """Verifica se o extrato tem as colunas necessárias para a separação,
    sem gerar planilha nenhuma. Usado pela interface para decidir se pergunta
    ou bloqueia antes de processar."""
    input_path = Path(input_path)
    rows = load_rows(input_path)
    return identify_header(rows)


def process_file(input_path: str | Path, output_dir: str | Path | None = None) -> ProcessResult:
    input_path = Path(input_path)
    if not input_path.exists():
        raise FileNotFoundError(f"Arquivo de entrada nao encontrado: {input_path}")

    rows = load_rows(input_path)
    header_match = identify_header(rows)
    if header_match.missing:
        raise ValueError(format_missing_columns_message(header_match.missing))
    header_idx, col_map = header_match.header_idx, header_match.col_map
    data_rows = ordenar_cronologicamente(rows, header_idx, col_map)
    groups = split_rows(data_rows, col_map)
    headers = output_headers(col_map)
    fields = output_fields(col_map)
    num_cols = len(headers)
    saldo_inicial, saldo_final = compute_balance_summary(data_rows, col_map)

    wb = Workbook()
    ws = wb.active
    ws.title = "Extrato Stone"

    current_row = 1
    if saldo_inicial is not None:
        escrever_linha_info(
            ws, current_row, f"Saldo antes do primeiro lançamento do mês: {saldo_inicial}", num_cols
        )
        current_row += 1
    if saldo_final is not None:
        escrever_linha_info(
            ws, current_row, f"Saldo final após o último lançamento do mês: {saldo_final}", num_cols
        )
        current_row += 1
    if saldo_inicial is not None or saldo_final is not None:
        linha_em_branco(ws, current_row, num_cols)
        current_row += 1

    freeze_row = current_row + 2

    current_row = escrever_secao(
        ws,
        current_row,
        "CRÉDITOS",
        ["Resgate da Reserva Stone", "Créditos comuns"],
        groups,
        headers,
        fields,
        num_cols,
    )

    linha_em_branco(ws, current_row, num_cols)
    current_row += 1
    linha_em_branco(ws, current_row, num_cols)
    current_row += 1

    current_row = escrever_secao(
        ws,
        current_row,
        "DÉBITOS",
        ["Aplicação na Reserva Stone", "Débitos comuns"],
        groups,
        headers,
        fields,
        num_cols,
    )

    if groups["Outros"]:
        linha_em_branco(ws, current_row, num_cols)
        current_row += 1
        current_row = escrever_secao(
            ws, current_row, "OUTROS", ["Outros"], groups, headers, fields, num_cols
        )

    ajustar_colunas(ws, fields)
    ws.freeze_panes = f"A{freeze_row}"

    output_root = Path(output_dir) if output_dir else None
    output_path = default_output_path(input_path, output_root)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()

    credit_rows = len(groups["Resgate da Reserva Stone"]) + len(groups["Créditos comuns"])
    debit_rows = len(groups["Aplicação na Reserva Stone"]) + len(groups["Débitos comuns"])
    return ProcessResult(
        output_path=output_path,
        total_rows=credit_rows + debit_rows + len(groups["Outros"]),
        entradas=credit_rows,
        saidas=debit_rows,
        banco=BANK_STONE,
    )
